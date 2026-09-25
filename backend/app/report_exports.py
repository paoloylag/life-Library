from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.charts.linecharts import HorizontalLineChart
from reportlab.graphics.charts.piecharts import Pie
from reportlab.graphics.shapes import Drawing
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

MAROON = "690F0D"
TEAL = "008C9B"
IVORY = "F2E8DC"


def safe_cell(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def add_sheet(book, name, headers, rows):
    sheet = book.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append([safe_cell(value) for value in row])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=MAROON)
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 27
    for column in sheet.columns:
        letter = get_column_letter(column[0].column)
        sheet.column_dimensions[letter].width = min(48, max(14, *(len(str(cell.value or "")) + 2 for cell in column)))
    return sheet


def export_excel(report: dict) -> bytes:
    book = Workbook()
    summary = book.active
    summary.title = "Executive Summary"
    summary.column_dimensions["A"].width = 35
    summary.column_dimensions["B"].width = 42
    summary.column_dimensions["C"].width = 40
    summary.merge_cells("A1:C1")
    summary["A1"] = "Life College | Library Attendance"
    summary["A1"].fill = PatternFill("solid", fgColor=MAROON)
    summary["A1"].font = Font(name="Aptos Display", size=17, bold=True, color="FFFFFF")
    summary.row_dimensions[1].height = 32

    def section(title):
        summary.append([])
        summary.append([title])
        row = summary.max_row
        summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        cell = summary.cell(row, 1)
        cell.fill = PatternFill("solid", fgColor=MAROON)
        cell.font = Font(bold=True, color="FFFFFF")
        summary.row_dimensions[row].height = 23

    filters = report["filters"]
    stats = report["summary"]
    section("Complete Report Details")
    for label, value in (
        ("Academic Year", filters["academic_year"]), ("Semester", filters["semester"]),
        ("Reporting Period", f'{report["period_start"] or "-"} to {report["period_end"] or "-"}'),
        ("Trend Grouping", filters["grouping"]), ("Generated (Asia/Manila)", report["generated_at"]),
        ("Calendar", report["calendar_note"]), ("Profile Data", report["profile_note"]),
    ):
        summary.append([label, safe_cell(value)])
    section("Summary Statistics")
    for key, label in (
        ("total_visits", "Total Library Visits"), ("unique_users", "Unique Users"),
        ("return_visits", "Return Visits"), ("returning_users", "Returning Users"),
        ("average_per_open_day", "Average Visits per Open Day"),
        ("average_per_week", "Average Visits per Week"),
        ("average_per_month", "Average Visits per Month"),
        ("average_per_user", "Average Visits per User"),
    ):
        summary.append([label, stats[key]])
    summary.append(["Open Days", stats["open_days"]])
    summary.append(["Peak Day", str(stats["peak_day"] or "-")])
    summary.append(["Peak Hour", str(stats["peak_hour"] or "-")])
    section("User Categories")
    for item in report["breakdowns"]["categories"]:
        summary.append([safe_cell(item["label"]), item["value"]])
    section("Student Breakdown")
    for dimension, key in (("Program / Course", "programs"), ("Year Level", "year_levels"), ("Section", "sections")):
        for item in report["breakdowns"][key]:
            summary.append([dimension, safe_cell(item["label"]), item["value"]])
    for row in summary.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    data = report["breakdowns"]
    sheets = {}
    for name, key in (
        ("Attendance Trend", "trend"), ("Daily Traffic", "daily"),
        ("Weekly Traffic", "weekly"), ("Monthly Traffic", "monthly"),
        ("Peak Hours", "hours"), ("Weekdays", "weekdays"),
        ("User Categories", "categories"), ("Programs", "programs"),
        ("Year Levels", "year_levels"), ("Sections", "sections"),
        ("Year and Section", "year_sections"), ("Semester Comparison", "semesters"),
    ):
        sheets[key] = add_sheet(book, name, ["Dimension", "Visits"], [(item["label"], item["value"]) for item in data[key]])
    columns = [
        ("date", "Date"), ("check_in_time", "Check-in Time"), ("name", "Name"),
        ("user_number", "User Number"), ("category", "User Category"),
        ("academic_year", "Academic Year"), ("semester", "Semester"),
        ("program", "Program"), ("year_level", "Year Level"),
        ("section", "Section"), ("department", "Department"),
        ("organization", "Organization"), ("source", "Source"),
        ("purpose", "Purpose"), ("note", "Note"),
    ]
    add_sheet(book, "Check-ins", [label for _, label in columns], [[row[key] for key, _ in columns] for row in report["records"]])
    for key, chart_type, anchor in (
        ("trend", LineChart, "E3"),
        ("categories", PieChart, "E20"),
        ("programs", BarChart, "E37"),
        ("year_levels", BarChart, "E54"),
        ("sections", BarChart, "E71"),
    ):
        sheet = sheets[key]
        if sheet.max_row <= 1:
            continue
        chart = chart_type()
        chart.title = {
            "trend": "Attendance Trend", "categories": "Visits by User Category",
            "programs": "Student Visits by Program", "year_levels": "Student Visits by Year Level",
            "sections": "Student Visits by Section",
        }[key]
        chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=sheet.max_row), titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=sheet.max_row))
        chart.width, chart.height = 17, 8
        summary.add_chart(chart, anchor)
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def export_pdf(report: dict) -> bytes:
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(letter), leftMargin=28, rightMargin=28)
    styles = getSampleStyleSheet()
    story = [Paragraph("Life College | Library Attendance", styles["Title"]), Spacer(1, 10)]

    def table(headers, rows, widths=None):
        body = [[str(value) for value in headers]] + [[str(value) for value in row] for row in rows]
        if len(body) == 1:
            body.append(["No records"] + [""] * (len(headers) - 1))
        element = Table(body, colWidths=widths, repeatRows=1, hAlign="CENTER")
        element.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + MAROON)),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F5F2")]),
            ("FONTSIZE", (0, 0), (-1, -1), 7), ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        story.extend([element, Spacer(1, 12)])

    def heading(title):
        story.append(Paragraph(title, styles["Heading2"]))

    def chart(items, kind):
        shown = items[:12]
        if not shown or not any(item["value"] for item in shown):
            return
        drawing = Drawing(440, 190)
        drawing.hAlign = "CENTER"
        if kind == "pie":
            graph = Pie()
            graph.x, graph.y, graph.width, graph.height = 120, 15, 170, 170
            graph.data = [item["value"] for item in shown]
            graph.labels = [item["label"][:14] for item in shown]
        elif kind == "line":
            graph = HorizontalLineChart()
            graph.x, graph.y, graph.width, graph.height = 45, 35, 360, 130
            graph.data = [[item["value"] for item in shown]]
            graph.categoryAxis.categoryNames = [item["label"][-7:] for item in shown]
        else:
            graph = VerticalBarChart()
            graph.x, graph.y, graph.width, graph.height = 45, 35, 360, 130
            graph.data = [[item["value"] for item in shown]]
            graph.categoryAxis.categoryNames = [item["label"][:12] for item in shown]
        drawing.add(graph)
        story.extend([drawing, Spacer(1, 8)])

    filters, stats, data = report["filters"], report["summary"], report["breakdowns"]
    heading("Complete Report Details")
    table(["Detail", "Value"], [
        ("Academic Year", filters["academic_year"]), ("Semester", filters["semester"]),
        ("Reporting Period", f'{report["period_start"] or "-"} to {report["period_end"] or "-"}'),
        ("Trend Grouping", filters["grouping"]), ("Generated (Asia/Manila)", report["generated_at"]),
    ], [180, 400])
    heading("Summary Statistics")
    table(["Measure", "Value"], [(key.replace("_", " ").title(), value) for key, value in stats.items()], [180, 400])
    heading("Visits by User Category")
    chart(data["categories"], "pie")
    table(["Category", "Visits"], [(row["label"], row["value"]) for row in data["categories"]], [250, 100])
    for title, key in (("Attendance Trend", "trend"), ("Student Visits by Program", "programs"),
                       ("Student Visits by Year Level", "year_levels"), ("Student Visits by Section", "sections"),
                       ("Semester Comparison", "semesters"), ("Peak Hours", "hours")):
        heading(title)
        if key in ("trend", "programs", "year_levels", "sections"):
            chart(data[key], "line" if key == "trend" else "bar")
        table(["Dimension", "Visits"], [(row["label"], row["value"]) for row in data[key]], [250, 100])
    story.append(Paragraph(report["calendar_note"], styles["Normal"]))
    story.append(Paragraph(report["profile_note"], styles["Normal"]))
    heading("Filtered Check-ins")
    fields = ("date", "name", "user_number", "category", "program", "year_level", "section", "source")
    table([field.replace("_", " ").title() for field in fields],
          [[row[field] for field in fields] for row in report["records"]],
          [65, 105, 82, 95, 125, 70, 70, 50])
    doc.build(story)
    return output.getvalue()
